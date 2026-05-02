from __future__ import annotations

import datetime
from pathlib import Path

import pytest

from retech_part2.ingestion.bilan.dates import combine_date_time, fix_bilan_date

SAMPLE_FILE = Path(__file__).parent.parent / "data" / "samples" / "avril-report1_2442026.xlsx"

# Spec, M4: the "Date" label is on row 10 of the source sheet. Hardcoding here
# is fine — M4 will discover it dynamically once the parser is built.
DATE_ROW = 10
FIRST_DATA_COLUMN = 5


# ---------------------------------------------------------------------------
# Cases the spec calls out explicitly.
# ---------------------------------------------------------------------------

def test_april_1_was_corrupted_to_jan_4() -> None:
    assert fix_bilan_date(datetime.datetime(2025, 1, 4)) == datetime.date(2025, 4, 1)


def test_april_12_was_corrupted_to_dec_4() -> None:
    assert fix_bilan_date(datetime.datetime(2025, 12, 4)) == datetime.date(2025, 4, 12)


def test_april_13_stayed_as_string() -> None:
    assert fix_bilan_date("13/04/2025") == datetime.date(2025, 4, 13)


# ---------------------------------------------------------------------------
# Edge cases that protect the parser from itself.
# ---------------------------------------------------------------------------

def test_none_returns_none() -> None:
    assert fix_bilan_date(None) is None


def test_empty_string_returns_none() -> None:
    assert fix_bilan_date("   ") is None


def test_garbage_string_returns_none() -> None:
    assert fix_bilan_date("not-a-date") is None


def test_iso_string_is_accepted_as_is() -> None:
    assert fix_bilan_date("2025-04-15") == datetime.date(2025, 4, 15)


def test_already_a_date_passes_through_unchanged() -> None:
    d = datetime.date(2025, 4, 7)
    assert fix_bilan_date(d) == d


def test_invalid_swap_falls_back_to_as_stored() -> None:
    # If swap produces an invalid date (here: day=15, month=2 → would become
    # month=15) we must NOT crash. Fall back to the cell as stored.
    raw = datetime.datetime(2025, 2, 15, 10, 0)
    assert fix_bilan_date(raw) == datetime.date(2025, 2, 15)


def test_combine_date_time_with_time() -> None:
    assert combine_date_time(datetime.date(2025, 4, 1), datetime.time(13, 27)) == datetime.datetime(2025, 4, 1, 13, 27)


def test_combine_date_time_without_time_uses_midnight() -> None:
    assert combine_date_time(datetime.date(2025, 4, 1), None) == datetime.datetime(2025, 4, 1, 0, 0)


def test_combine_date_time_without_date_returns_none() -> None:
    assert combine_date_time(None, datetime.time(13, 27)) is None


# ---------------------------------------------------------------------------
# Integration: the real April-2025 sample must yield only April dates after
# normalisation. Skipped when the file isn't present locally.
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not SAMPLE_FILE.exists(), reason=f"sample file not present at {SAMPLE_FILE}")
def test_real_file_yields_only_april_2025() -> None:
    import openpyxl  # local import keeps the rest of the module test-tool-light

    wb = openpyxl.load_workbook(SAMPLE_FILE, read_only=True, data_only=True)
    ws = wb.active

    fixed_dates: list[datetime.date] = []
    raw_strings_seen = 0
    raw_datetimes_seen = 0

    for col in range(FIRST_DATA_COLUMN, ws.max_column + 1):
        raw = ws.cell(row=DATE_ROW, column=col).value
        if raw is None:
            continue
        if isinstance(raw, str):
            raw_strings_seen += 1
        elif isinstance(raw, datetime.datetime):
            raw_datetimes_seen += 1
        d = fix_bilan_date(raw)
        if d is not None:
            fixed_dates.append(d)

    wb.close()

    assert fixed_dates, "expected at least one fixed date from the sample sheet"
    assert raw_datetimes_seen > 0, "expected days 1..12 to come through as datetime cells"
    assert raw_strings_seen > 0, "expected days 13..20 to come through as string cells"

    months = {d.month for d in fixed_dates}
    assert months == {4}, f"expected only April after normalisation, got months={months}"

    days = {d.day for d in fixed_dates}
    assert days <= set(range(1, 21)), f"expected days within 1..20, got {sorted(days)}"

    years = {d.year for d in fixed_dates}
    assert years == {2025}, f"expected year 2025 only, got {years}"
