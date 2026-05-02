"""Structural parser for BILAN Excel files (M4).

In-memory only — no DB writes. Output is a `ParsedBilan` with one
`ParsedReading` per (timestamp × parameter row) where the cell holds a value.

Handles the landmines documented in the project brief:
  * Date-cell locale swap (delegated to `bilan/dates.py`).
  * "Date" / "Heure" header rows discovered dynamically in column B.
  * Merged-cell categories in column A — forward-filled from the top-left of
    each merged range (pandas ffill cannot do this).
  * Whitespace anomalies in raw labels — collapsed before storage.
  * ~927 "gap columns" with values but no date/time — timestamps synthesized by
    linear interpolation between the nearest neighbouring valid timestamps,
    with a flag.
  * Empty parameter rows (e.g. row 13 has zero values) — recorded in
    `parameter_summary` with count 0; no readings emitted.
  * Out-of-order timestamps within a single day — readings remain in the parser
    output in COLUMN ORDER (the true logging sequence). The `monotonic_inversion`
    flag is applied later in M6 once the YAML mapping identifies which metrics
    are cumulative meters.
"""
from __future__ import annotations

import datetime as dt
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from retech_part2.ingestion.bilan.dates import combine_date_time, fix_bilan_date
from retech_part2.utils.hashing import sha256_file

DATA_FIRST_COLUMN = 5
HEADER_LABEL_DATE = "date"
HEADER_SCAN_MAX_ROW = 30
TIME_ROW_SCAN_RANGE = 5  # rows after date_row to scan for the time row
DEFAULT_LOG_INTERVAL = dt.timedelta(minutes=10)
WHITESPACE_RE = re.compile(r"\s+")


@dataclass
class ParsedReading:
    time: dt.datetime
    raw_label: str
    category: str
    value: float
    timestamp_synthetic: bool = False
    data_quality_flags: list[str] = field(default_factory=list)


@dataclass
class ParsedBilan:
    file_hash: str
    filename: str
    date_range: tuple[dt.date, dt.date]
    readings: list[ParsedReading]
    parameter_summary: dict[str, int]
    parse_warnings: list[str]


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _normalize_label(s: Any) -> str | None:
    if s is None:
        return None
    if not isinstance(s, str):
        s = str(s)
    s = WHITESPACE_RE.sub(" ", s).strip()
    return s or None


def _find_label_row(ws: Worksheet, *, col: int, label: str) -> int | None:
    """First row in `col` whose normalised text equals `label` (case-insensitive)."""
    target = label.casefold()
    for row in range(1, HEADER_SCAN_MAX_ROW + 1):
        v = _normalize_label(ws.cell(row=row, column=col).value)
        if v and v.casefold() == target:
            return row
    return None


def _find_time_row(ws: Worksheet, *, after_row: int) -> int | None:
    """Find the row that holds per-column time-of-day values.

    The BILAN sheet's column-B label for this row is verbose (e.g.
    "Heure de l'inspection périodique"), so we identify it structurally:
    the first row after the date row whose first data cell is a
    ``datetime.time`` (or ``datetime.datetime``).
    """
    for row in range(after_row + 1, after_row + 1 + TIME_ROW_SCAN_RANGE):
        v = ws.cell(row=row, column=DATA_FIRST_COLUMN).value
        if isinstance(v, dt.time) and not isinstance(v, dt.datetime):
            return row
        if isinstance(v, dt.datetime):
            return row
    return None


def _build_col_a_map(ws: Worksheet) -> dict[int, str | None]:
    """Map row → column-A category, with merged-range forward-fill.

    pandas `ffill()` only sees the top-left value of a merged range, so any
    parameter row whose category is hidden inside a merged span would lose its
    category. We resolve merges explicitly here.
    """
    result: dict[int, str | None] = {}
    for row_idx in range(1, ws.max_row + 1):
        result[row_idx] = _normalize_label(ws.cell(row=row_idx, column=1).value)
    for mr in ws.merged_cells.ranges:
        if not (mr.min_col <= 1 <= mr.max_col):
            continue
        top_left = _normalize_label(ws.cell(row=mr.min_row, column=mr.min_col).value)
        for row_idx in range(mr.min_row, mr.max_row + 1):
            result[row_idx] = top_left
    return result


def _coerce_time(raw: Any) -> dt.time | None:
    """Cell value can be a `time`, a `datetime`, a string like '01:27', or None."""
    if raw is None:
        return None
    # `datetime` is a subclass of `date` (not `time`); check the time subclass first.
    if isinstance(raw, dt.time) and not isinstance(raw, dt.datetime):
        return raw
    if isinstance(raw, dt.datetime):
        return raw.time()
    if isinstance(raw, str):
        s = raw.strip()
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                return dt.datetime.strptime(s, fmt).time()
            except ValueError:
                continue
    return None


def _build_timestamp_index(
    ws: Worksheet,
    *,
    date_row: int,
    time_row: int,
    last_col: int,
) -> tuple[list[dt.datetime | None], set[int]]:
    """Two-pass timestamp construction.

    Returns (timestamps_per_data_column, set_of_synthesised_offsets).
    Index = column - DATA_FIRST_COLUMN.
    """
    cols = list(range(DATA_FIRST_COLUMN, last_col + 1))
    raw_ts: list[dt.datetime | None] = []
    for col in cols:
        d = fix_bilan_date(ws.cell(row=date_row, column=col).value)
        t = _coerce_time(ws.cell(row=time_row, column=col).value)
        raw_ts.append(combine_date_time(d, t))

    ts: list[dt.datetime | None] = list(raw_ts)
    synthetic: set[int] = set()
    valid_indices = [i for i, t in enumerate(raw_ts) if t is not None]
    if not valid_indices:
        return ts, synthetic

    # Quick lookups for nearest-neighbour search
    valid_set_sorted = valid_indices  # already in ascending order
    ptr_next = 0  # index into valid_set_sorted

    for i in range(len(ts)):
        if ts[i] is not None:
            continue
        # Advance ptr_next so valid_set_sorted[ptr_next] is the first valid index > i
        while ptr_next < len(valid_set_sorted) and valid_set_sorted[ptr_next] <= i:
            ptr_next += 1
        next_i = valid_set_sorted[ptr_next] if ptr_next < len(valid_set_sorted) else None
        prev_i = valid_set_sorted[ptr_next - 1] if ptr_next > 0 else None

        if prev_i is not None and next_i is not None:
            span = next_i - prev_i
            offset = i - prev_i
            delta = (raw_ts[next_i] - raw_ts[prev_i]) / span
            ts[i] = raw_ts[prev_i] + delta * offset
        elif prev_i is not None:
            ts[i] = raw_ts[prev_i] + DEFAULT_LOG_INTERVAL * (i - prev_i)
        elif next_i is not None:
            ts[i] = raw_ts[next_i] - DEFAULT_LOG_INTERVAL * (next_i - i)
        synthetic.add(i)

    return ts, synthetic


# ---------------------------------------------------------------------------
# main entry point
# ---------------------------------------------------------------------------

def parse_bilan(path: Path) -> ParsedBilan:
    file_hash = sha256_file(path)
    warnings: list[str] = []

    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        ws = wb.active
        if ws is None:
            raise ValueError(f"workbook {path.name} has no active sheet")

        date_row = _find_label_row(ws, col=2, label=HEADER_LABEL_DATE)
        if date_row is None:
            raise ValueError("could not find 'Date' label in column B")
        time_row = _find_time_row(ws, after_row=date_row)
        if time_row is None:
            raise ValueError(
                f"could not find time row after date_row={date_row} "
                f"(scanned {TIME_ROW_SCAN_RANGE} rows for time-typed data cells)"
            )

        first_param_row = max(date_row, time_row) + 1
        last_col = ws.max_column

        col_a_map = _build_col_a_map(ws)
        ts_by_offset, synthetic_offsets = _build_timestamp_index(
            ws, date_row=date_row, time_row=time_row, last_col=last_col
        )

        readings: list[ParsedReading] = []
        param_summary: dict[str, int] = {}

        for param_row in range(first_param_row, ws.max_row + 1):
            raw_label = _normalize_label(ws.cell(row=param_row, column=2).value)
            if not raw_label:
                continue

            category = col_a_map.get(param_row) or ""
            param_summary.setdefault(raw_label, 0)

            for col_offset, ts in enumerate(ts_by_offset):
                if ts is None:
                    continue
                col = col_offset + DATA_FIRST_COLUMN
                cell_value = ws.cell(row=param_row, column=col).value
                if cell_value is None:
                    continue
                try:
                    val = float(cell_value)
                except (TypeError, ValueError):
                    warnings.append(
                        f"non-numeric value at row={param_row} col={col}: {cell_value!r}"
                    )
                    continue

                readings.append(
                    ParsedReading(
                        time=ts,
                        raw_label=raw_label,
                        category=category,
                        value=val,
                        timestamp_synthetic=col_offset in synthetic_offsets,
                    )
                )
                param_summary[raw_label] += 1
    finally:
        wb.close()

    if not readings:
        raise ValueError(f"no readings parsed from {path.name}")

    times_only = [r.time for r in readings]
    date_range = (min(times_only).date(), max(times_only).date())

    return ParsedBilan(
        file_hash=file_hash,
        filename=path.name,
        date_range=date_range,
        readings=readings,
        parameter_summary=param_summary,
        parse_warnings=warnings,
    )


# ---------------------------------------------------------------------------
# CLI: python -m retech_part2.ingestion.bilan.parser path/to/file.xlsx
# ---------------------------------------------------------------------------

def _print_summary(p: ParsedBilan) -> None:
    label_w = max((len(k) for k in p.parameter_summary), default=20)
    label_w = min(label_w, 60)
    print(f"\nFile     : {p.filename}")
    print(f"Hash     : {p.file_hash[:16]}…")
    print(f"Range    : {p.date_range[0]} → {p.date_range[1]}")
    print(f"Readings : {len(p.readings):,}")
    n_synth = sum(1 for r in p.readings if r.timestamp_synthetic)
    print(f"Synthetic timestamps in readings : {n_synth:,} ({n_synth / len(p.readings):.1%})")
    print(f"Distinct parameters : {len(p.parameter_summary)}")
    print(f"Parse warnings      : {len(p.parse_warnings)}\n")

    print(f"{'raw_label'.ljust(label_w)}  {'count':>8}")
    print(f"{'-' * label_w}  {'-' * 8}")
    for label, count in sorted(p.parameter_summary.items(), key=lambda kv: (-kv[1], kv[0])):
        display = label if len(label) <= label_w else label[: label_w - 1] + "…"
        marker = "  (empty)" if count == 0 else ""
        print(f"{display.ljust(label_w)}  {count:>8,}{marker}")

    # Per-day synthesis histogram
    by_day = Counter(r.time.date() for r in p.readings if r.timestamp_synthetic)
    if by_day:
        print(f"\nSynthesised timestamps by day:")
        for day in sorted(by_day):
            print(f"  {day}: {by_day[day]:,}")

    if p.parse_warnings:
        print(f"\nFirst 5 warnings:")
        for w in p.parse_warnings[:5]:
            print(f"  - {w}")


def main(argv: list[str] | None = None) -> int:
    args = argv if argv is not None else sys.argv[1:]
    if len(args) != 1:
        print("usage: python -m retech_part2.ingestion.bilan.parser <path-to-xlsx>", file=sys.stderr)
        return 2
    path = Path(args[0]).resolve()
    if not path.exists():
        print(f"file not found: {path}", file=sys.stderr)
        return 1
    parsed = parse_bilan(path)
    _print_summary(parsed)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
